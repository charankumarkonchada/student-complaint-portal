import io
import datetime
from flask import Blueprint, redirect, url_for, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import backend.config as config
from backend.database.db import get_db_connection
from backend.services.auth_service import admin_required

export_reports_bp = Blueprint("export_reports", __name__)


def normalize_excel_date(val):
    """
    Safely normalizes various date/time representations (datetime, date, timestamp,
    ISO/custom strings) into a timezone-naive datetime.datetime for Excel/openpyxl.
    """
    if val is None:
        return ""
    if isinstance(val, datetime.datetime):
        if val.tzinfo is not None:
            val = val.astimezone().replace(tzinfo=None)
        return val
    if isinstance(val, datetime.date):
        return datetime.datetime(val.year, val.month, val.day)
    if isinstance(val, (int, float)):
        try:
            return datetime.datetime.fromtimestamp(val)
        except Exception:
            return str(val)
    if isinstance(val, str):
        val_str = val.strip()
        if not val_str:
            return ""
        for fmt in (
            "%Y-%m-%d %H:%M:%S.%f",
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%dT%H:%M:%S.%f",
            "%Y-%m-%dT%H:%M:%S",
            "%Y-%m-%d",
            "%d-%m-%Y %H:%M:%S",
            "%d-%m-%Y",
        ):
            try:
                return datetime.datetime.strptime(val_str, fmt)
            except ValueError:
                pass
        try:
            return datetime.datetime.fromisoformat(val_str.replace("Z", "+00:00")).replace(tzinfo=None)
        except Exception:
            return val_str
    return val

@export_reports_bp.route("/export_pdf")
def export_pdf():
    if not admin_required():
        return redirect(url_for("admin_login"))

    conn = get_db_connection()
    rows = conn.execute(
        """
        SELECT complaints.*, students.name, students.id_no
        FROM complaints
        JOIN students ON complaints.student_id=students.id
        ORDER BY complaints.created_at DESC
        """
    ).fetchall()
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib import colors

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.setTitle(f"{config.COLLEGE_NAME} Complaint Report")

    width, height = A4
    margin = 40
    printable_width = width - (2 * margin)

    title_style = ParagraphStyle(
        name="ReportTitle",
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#0f172a")
    )

    title_text = f"<b>{config.COLLEGE_NAME}</b><br/><font size='11' color='#475569'>Hostel Complaint Report</font>"
    title_p = Paragraph(title_text, title_style)
    _, title_h = title_p.wrap(printable_width, height)

    y = height - margin - title_h
    title_p.drawOn(pdf, margin, y)

    y -= 10
    pdf.setStrokeColor(colors.HexColor("#cbd5e1"))
    pdf.setLineWidth(1)
    pdf.line(margin, y, width - margin, y)
    y -= 20

    pdf.setFont("Helvetica", 9)
    pdf.setFillColor(colors.HexColor("#1e293b"))

    for c in rows:
        line = (
            f"#{c['id']} | "
            f"{c['name']} | "
            f"{dict(c).get('id_no', '')} | "
            f"{c['category']} | "
            f"{c['priority']} | "
            f"{c['status']} | "
            f"AI: {c['ai_resolution_days'] or '-'}d"
        )
        pdf.drawString(margin, y, line[:115])
        y -= 16
        if y < 45:
            pdf.showPage()
            y = height - margin
            pdf.setFont("Helvetica", 9)
            pdf.setFillColor(colors.HexColor("#1e293b"))

    pdf.save()
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="RGUKT_Ongole_Complaint_Report.pdf",
        mimetype="application/pdf"
    )

@export_reports_bp.route("/export_excel")
def export_excel():
    if not admin_required():
        return redirect(url_for("admin_login"))

    conn = get_db_connection()
    rows = conn.execute(
        """
        SELECT complaints.*, students.name, students.id_no, students.hostel
        FROM complaints
        JOIN students ON complaints.student_id=students.id
        ORDER BY complaints.created_at DESC
        """
    ).fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Complaints"
    ws.freeze_panes = "A2"

    headers = [
        "ID", "Student", "ID No", "Hostel", "Category", "Title", "Priority",
        "Status", "Assigned To", "AI Category", "AI Category Confidence",
        "AI Priority", "AI Priority Confidence", "AI Resolution Days",
        "Duplicate ID", "Similarity %", "Date"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", start_color="123B5D", end_color="123B5D")
    header_align = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    date_col_idx = 17
    for r_idx, c in enumerate(rows, start=2):
        created_val = normalize_excel_date(c["created_at"])
        ws.append([
            c["id"],
            c["name"],
            dict(c).get("id_no", ""),
            dict(c).get("hostel", ""),
            c["category"],
            c["title"],
            c["priority"],
            c["status"],
            c["assigned_to"],
            c["ai_category"],
            c["ai_category_confidence"],
            c["ai_priority"],
            c["ai_priority_confidence"],
            c["ai_resolution_days"],
            c["ai_duplicate_id"],
            c["ai_duplicate_similarity"],
            created_val
        ])
        if isinstance(created_val, (datetime.datetime, datetime.date)):
            date_cell = ws.cell(row=r_idx, column=date_col_idx)
            date_cell.number_format = "dd-mm-yyyy hh:mm AM/PM"
            date_cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = cell.value
            if val is None:
                continue
            if isinstance(val, (datetime.datetime, datetime.date)):
                val_str = val.strftime("%d-%m-%Y %I:%M %p")
            else:
                val_str = str(val)
            max_len = max(max_len, len(val_str))

        if col_letter == "Q":  # Date column
            ws.column_dimensions[col_letter].width = max(max_len + 3, 22)
        elif col_letter == "F":  # Title column
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 15), 40)
            for cell in col[1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        elif col_letter == "A":  # ID column
            ws.column_dimensions[col_letter].width = max(max_len + 3, 8)
        else:
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 32)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="RGUKT_Ongole_Complaint_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
