import io
import os
import sys
import unittest
import datetime
import openpyxl

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

os.environ["TESTING"] = "1"
import backend.config as config
from app import create_app
from backend.database.db import get_db_connection
from backend.database.queries import init_database
from backend.routes.charankumar.export_reports import normalize_excel_date

TEST_DB_PATH = os.path.join(BASE_DIR, "test_export_excel.db")


class TestExcelExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if os.path.exists(TEST_DB_PATH):
            try:
                os.remove(TEST_DB_PATH)
            except Exception:
                pass

        config.DATABASE_URL = ""
        config.DATABASE = TEST_DB_PATH
        config.SECRET_KEY = "test-secret-key-export"

        init_database()
        cls.app = create_app()
        cls.app.config.update({
            "TESTING": True,
            "WTF_CSRF_ENABLED": False,
        })
        cls.client = cls.app.test_client()

        # Seed student and complaint
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute(
            """
            INSERT OR REPLACE INTO students (id, name, id_no, email, phone, hostel, room_no, password)
            VALUES (1, 'Test Student', 'O220001', 'o220001@rguktong.ac.in', '9876543210', 'BH-1', '101', 'hash')
            """
        )
        cur.execute(
            """
            INSERT INTO complaints (
                id, student_id, category, title, description, priority, status,
                assigned_to, ai_category, ai_category_confidence, ai_priority,
                ai_priority_confidence, ai_resolution_days, ai_duplicate_id,
                ai_duplicate_similarity, created_at
            )
            VALUES (
                101, 1, 'Electrical', 'Ceiling Fan Not Working', 'Speed regulator damaged', 'High', 'In Progress',
                'Electrician Ramu', 'Electrical', 0.95, 'High', 0.90, 2, NULL, NULL, '2026-09-29 17:36:37'
            )
            """
        )
        conn.commit()
        conn.close()

    @classmethod
    def tearDownClass(cls):
        if os.path.exists(TEST_DB_PATH):
            try:
                os.remove(TEST_DB_PATH)
            except Exception:
                pass

    def test_normalize_excel_date_datetime(self):
        dt = datetime.datetime(2026, 9, 29, 17, 36, 37)
        self.assertEqual(normalize_excel_date(dt), dt)

    def test_normalize_excel_date_timezone_aware(self):
        tz = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
        dt_aware = datetime.datetime(2026, 9, 29, 17, 36, 37, tzinfo=tz)
        res = normalize_excel_date(dt_aware)
        self.assertIsNone(res.tzinfo)

    def test_normalize_excel_date_date_object(self):
        d = datetime.date(2026, 9, 29)
        res = normalize_excel_date(d)
        self.assertIsInstance(res, datetime.datetime)
        self.assertEqual(res.year, 2026)
        self.assertEqual(res.month, 9)
        self.assertEqual(res.day, 29)

    def test_normalize_excel_date_strings(self):
        s1 = "2026-09-29 17:36:37"
        res1 = normalize_excel_date(s1)
        self.assertIsInstance(res1, datetime.datetime)
        self.assertEqual(res1.year, 2026)
        self.assertEqual(res1.minute, 36)

        s2 = "2026-09-29T17:36:37"
        res2 = normalize_excel_date(s2)
        self.assertIsInstance(res2, datetime.datetime)

        s3 = "29-09-2026 17:36:37"
        res3 = normalize_excel_date(s3)
        self.assertIsInstance(res3, datetime.datetime)

    def test_normalize_excel_date_none_and_empty(self):
        self.assertEqual(normalize_excel_date(None), "")
        self.assertEqual(normalize_excel_date(""), "")
        self.assertEqual(normalize_excel_date("   "), "")

    def test_export_excel_requires_admin(self):
        res = self.client.get("/export_excel")
        self.assertEqual(res.status_code, 302)
        self.assertIn("/admin_login", res.headers.get("Location", ""))

    def test_export_excel_success(self):
        with self.client.session_transaction() as sess:
            sess["admin"] = "admin"

        res = self.client.get("/export_excel")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res.headers.get("Content-Type"),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        self.assertIn("RGUKT_Ongole_Complaint_Report.xlsx", res.headers.get("Content-Disposition", ""))

        # Parse workbook
        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        self.assertIn("Complaints", wb.sheetnames)
        ws = wb["Complaints"]

        # Check freeze panes
        self.assertEqual(ws.freeze_panes, "A2")

        # Check headers
        expected_headers = [
            "ID", "Student", "ID No", "Hostel", "Category", "Title", "Priority",
            "Status", "Assigned To", "AI Category", "AI Category Confidence",
            "AI Priority", "AI Priority Confidence", "AI Resolution Days",
            "Duplicate ID", "Similarity %", "Date"
        ]
        header_vals = [ws.cell(row=1, column=i).value for i in range(1, 18)]
        self.assertEqual(header_vals, expected_headers)

        # Check header style
        for col_idx in range(1, 18):
            cell = ws.cell(row=1, column=col_idx)
            self.assertTrue(cell.font.bold)
            self.assertEqual(cell.font.color.rgb, "00FFFFFF")
            self.assertEqual(cell.fill.start_color.rgb, "00123B5D")

        # Check row 2 (data row)
        self.assertEqual(ws.cell(row=2, column=1).value, 101)
        self.assertEqual(ws.cell(row=2, column=2).value, "Test Student")
        self.assertEqual(ws.cell(row=2, column=6).value, "Ceiling Fan Not Working")

        # Check Date column (Column 17 / Q)
        date_cell = ws.cell(row=2, column=17)
        self.assertIsInstance(date_cell.value, datetime.datetime)
        self.assertEqual(date_cell.number_format, "dd-mm-yyyy hh:mm AM/PM")
        self.assertEqual(date_cell.alignment.horizontal, "center")

        # Check column dimensions
        q_width = ws.column_dimensions["Q"].width
        self.assertIsNotNone(q_width)
        self.assertGreaterEqual(q_width, 22.0)

        f_width = ws.column_dimensions["F"].width
        self.assertIsNotNone(f_width)
        self.assertGreaterEqual(f_width, 15.0)
        self.assertLessEqual(f_width, 40.0)


if __name__ == "__main__":
    unittest.main()
